from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Define color scheme
RED = RGBColor(220, 20, 60)  # Crimson
WHITE = RGBColor(255, 255, 255)
DARK_GRAY = RGBColor(40, 40, 40)
LIGHT_GRAY = RGBColor(245, 245, 245)

def add_title_slide(prs, title, subtitle):
    """Add a title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RED
    
    # Title
    left = Inches(0.5)
    top = Inches(2.5)
    width = Inches(9)
    height = Inches(1.5)
    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    left = Inches(0.5)
    top = Inches(4.2)
    width = Inches(9)
    height = Inches(1.5)
    subtitle_box = slide.shapes.add_textbox(left, top, width, height)
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(32)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_points):
    """Add a content slide with bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    # Title
    left = Inches(0.5)
    top = Inches(0.4)
    width = Inches(9)
    height = Inches(0.8)
    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RED
    
    # Add red underline
    shape = slide.shapes.add_shape(1, Inches(0.5), Inches(1.3), Inches(9), Inches(0))
    shape.line.color.rgb = RED
    shape.line.width = Pt(3)
    
    # Content
    left = Inches(1)
    top = Inches(1.7)
    width = Inches(8)
    height = Inches(5)
    text_box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    
    for i, point in enumerate(content_points):
        if i > 0:
            text_frame.add_paragraph()
        p = text_frame.paragraphs[i]
        p.text = point
        p.font.size = Pt(24)
        p.font.color.rgb = DARK_GRAY
        p.level = 0
        p.space_before = Pt(10)
        p.space_after = Pt(10)
    
    return slide

# Slide 1: Title
add_title_slide(prs, "Blood Donation", "Web Application")

# Slide 2: Overview
add_content_slide(prs, "Project Overview", [
    "🩸 A full-fledged web application connecting blood donors with people in emergency need",
    "📍 Geolocation-based donor search and matching",
    "⏱️ Real-time communication between donors and requestors",
    "💡 Modern, responsive UI with Tailwind CSS",
    "🔐 Secure user authentication for donors and emergency requestors"
])

# Slide 3: Key Features
add_content_slide(prs, "Key Features", [
    "✅ User authentication (Donor & Emergency Requestor roles)",
    "✅ Geolocation-based donor search",
    "✅ Real-time messaging with Socket.io",
    "✅ Donor profile management and blood type filtering",
    "✅ Blood request creation and tracking",
    "✅ Responsive design for mobile & desktop"
])

# Slide 4: Technology Stack
add_content_slide(prs, "Technology Stack", [
    "🎨 Frontend: React 18, Tailwind CSS, Axios",
    "⚙️ Backend: Node.js, Express.js, MongoDB",
    "🔄 Real-time: Socket.io",
    "🗺️ Maps: Google Maps API / Mapbox",
    "📱 Responsive Design: Mobile First Approach"
])

# Slide 5: Architecture
add_content_slide(prs, "Application Architecture", [
    "/client - React frontend application",
    "/server - Express.js backend API with MongoDB",
    "Modular component structure for easy customization",
    "RESTful API endpoints for all operations",
    "Real-time updates via WebSocket connections"
])

# Slide 6: User Roles
add_content_slide(prs, "User Roles", [
    "👨‍⚕️ Donors: Register blood type, location, availability",
    "📋 Create and manage donor profiles",
    "⚡ Emergency Requestors: Search for nearby donors",
    "💬 Send requests and communicate with donors",
    "📊 Track blood request status in real-time"
])

# Slide 7: Core Pages
add_content_slide(prs, "Core Pages", [
    "🏠 Home - Landing page with overview",
    "🔐 Login/Register - Secure authentication",
    "📍 Find Donors - Search and filter donors by location & blood type",
    "🩸 Request Blood - Create and track blood requests",
    "👤 Profile - Manage personal and blood donation information",
    "📊 Dashboard - View statistics and activity"
])

# Slide 8: Database Schema
add_content_slide(prs, "Database Models", [
    "User Model: Email, password, name, blood type, location, role",
    "Blood Request Model: Request details, urgency, status, requester info",
    "Donation Model: Donation history, date, amount, recipient info",
    "All models include timestamps and reference relationships"
])

# Slide 9: Workflow
add_content_slide(prs, "Typical Workflow", [
    "1️⃣ Emergency requestor creates a blood request",
    "2️⃣ System searches nearby donors by geolocation",
    "3️⃣ Notifications sent to eligible donors",
    "4️⃣ Donors can accept or decline requests",
    "5️⃣ Real-time messaging between donor & requestor",
    "6️⃣ Request status tracked through completion"
])

# Slide 10: Security Features
add_content_slide(prs, "Security & Authentication", [
    "✔️ Password hashing and secure storage",
    "✔️ JWT-based authentication tokens",
    "✔️ Protected API endpoints with middleware",
    "✔️ Role-based access control",
    "✔️ Data validation and sanitization"
])

# Slide 11: Future Enhancements
add_content_slide(prs, "Future Enhancements", [
    "🏥 Hospital integration and blood inventory management",
    "📊 Advanced analytics and donor statistics",
    "🔔 Push notifications for mobile app",
    "🌍 Multi-language support",
    "⭐ Rating and review system",
    "🤝 Community donation drives and events"
])

# Slide 12: Getting Started
add_content_slide(prs, "Getting Started", [
    "1. Clone the repository",
    "2. Install dependencies: npm install (both client & server)",
    "3. Configure MongoDB connection",
    "4. Set environment variables (.env)",
    "5. Run: npm start (server) and npm start (client)",
    "6. Access at http://localhost:3000"
])

# Slide 13: Customization
add_content_slide(prs, "Customization Points", [
    "🎨 API endpoints modifiable in server routes",
    "🧩 UI components are modular and customizable",
    "📚 Database schema can be extended",
    "⚡ Real-time features scalable with Socket.io",
    "🔧 Configure Tailwind CSS for branding"
])

# Slide 14: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = RED

left = Inches(1)
top = Inches(2.5)
width = Inches(8)
height = Inches(2.5)
text_box = slide.shapes.add_textbox(left, top, width, height)
text_frame = text_box.text_frame
text_frame.word_wrap = True

p = text_frame.paragraphs[0]
p.text = "Save Lives Through Technology"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

text_frame.add_paragraph()
p = text_frame.paragraphs[1]
p.text = "Making Blood Donation Accessible & Easy"
p.font.size = Pt(28)
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(20)

# Save presentation
prs.save('Blood_Donation_Presentation.pptx')
print("✅ Presentation created successfully: Blood_Donation_Presentation.pptx")
